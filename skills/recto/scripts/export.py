#!/usr/bin/env python3
"""Export one semantic Markdown source to Recto Typst/PDF/editable Word.

Run doctor.py first; references/setup.md handles missing dependencies separately.
Pandoc >=3.9 and Typst >=0.15.1 are external executables. Export runs offline.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from hashlib import sha256
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from doctor import check, FONT_DIR

try:
    import docx
    import lxml.etree
except (ImportError, OSError) as error:
    raise SystemExit(f'{error}\nRun scripts/doctor.py, then follow references/setup.md.')

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'm': MATH_NS}
STYLES = {
    'reference': dict(size=12, title_size=16.5, line=18, margins=(12,12,9.7896,9.7896)),
    'gbt9704': dict(size=16, title_size=22, line=29.609, margins=(37,35,28,26)),
}


def run(args, **kwargs):
    result = subprocess.run(args, capture_output=True, check=False, **kwargs)
    error = result.stderr.decode() if isinstance(result.stderr, bytes) else result.stderr
    if result.returncode or error.strip():
        # Treat converter warnings as failures: missing equations must not silently ship.
        raise RuntimeError(f'{args[0]} failed or reported diagnostics:\n{error}')
    return result.stdout


def elem(name, **attrs):
    out = OxmlElement(name)
    for key, value in attrs.items(): out.set(qn(key.replace('_', ':')), str(value))
    return out


def font_style(style, east, latin, size, bold=False):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0,0,0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        fonts.attrib.pop(qn('w:'+key), None)
    for key, value in [('eastAsia',east),('ascii',latin),('hAnsi',latin),('cs',latin)]:
        fonts.set(qn('w:'+key), value)
    lang = rpr.find(qn('w:lang'))
    if lang is None: lang = elem('w:lang'); rpr.append(lang)
    lang.set(qn('w:eastAsia'), 'zh-CN')


def field(paragraph, code):
    r = paragraph.add_run()
    r._r.append(elem('w:fldChar', w_fldCharType='begin'))
    instruction = elem('w:instrText'); instruction.text = ' '+code+' '; r._r.append(instruction)
    r._r.append(elem('w:fldChar', w_fldCharType='separate'))
    t = elem('w:t'); t.text = '1'; r._r.append(t)
    r._r.append(elem('w:fldChar', w_fldCharType='end'))


def configure_section(section, profile, landscape=False):
    cfg = STYLES[profile]
    section.page_width = Mm(297 if landscape else 210)
    section.page_height = Mm(210 if landscape else 297)
    top,bottom,left,right = cfg['margins']
    if landscape: top,bottom,left,right = (28,26,37,35)
    for key, value in [('top',top),('bottom',bottom),('left',left),('right',right)]:
        setattr(section, key+'_margin', Mm(value))
    section.header_distance = Mm(0)
    section.footer_distance = Mm(21 if profile == 'gbt9704' and not landscape else 6)
    pg = section._sectPr.find(qn('w:pgSz'))
    if landscape: pg.set(qn('w:orient'), 'landscape')
    else: pg.attrib.pop(qn('w:orient'),None)


def make_reference(path, profile):
    path.write_bytes(run(['pandoc','--print-default-data-file=reference.docx']))
    doc = Document(path); cfg = STYLES[profile]; course = profile == 'reference'
    for name in ['Normal','Body Text','First Paragraph','Compact','Footnote Text','Caption',
                 'Image Caption','Table Caption','Title','Subtitle','Heading 1','Heading 2',
                 'Heading 3','Heading 4','Appendix Label','Appendix Title','Footer']:
        if name not in doc.styles: doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s = doc.styles[name]
        font_style(s, cfg['body'], cfg['latin'], cfg['size'])
        p = s.paragraph_format
        p.space_before = Pt(0); p.space_after = Pt(9 if course else 0)
        p.line_spacing = Pt(cfg['line']); p.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        p.first_line_indent = Pt(0 if course else 31.57)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if course else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.widow_control = True
        p.keep_with_next = False
    for name in ('Title', 'Appendix Title'):
        s = doc.styles[name]
        font_style(s,cfg['title'],cfg['latin'],cfg['title_size'],course or cfg['title_bold'])
        p=s.paragraph_format; p.first_line_indent=Pt(0)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if course else WD_ALIGN_PARAGRAPH.CENTER
        p.space_after=Pt(16 if course else cfg['line']); p.keep_with_next=True
    s=doc.styles['Subtitle']; font_style(s,cfg['body'],cfg['latin'],13.5,True)
    s.paragraph_format.first_line_indent=Pt(0); s.paragraph_format.keep_with_next=True
    if course:
        borders=elem('w:pBdr'); borders.append(elem('w:bottom',w_val='single',w_sz=4,w_space=10,w_color='000000'))
        s.element.get_or_add_pPr().append(borders)
    for level in range(1,5):
        s=doc.styles[f'Heading {level}']
        east=cfg['body'] if course or level>2 else (cfg['hei'] if level==1 else cfg['kai'])
        font_style(s,east,cfg['latin'],13.5 if course else 16,course)
        p=s.paragraph_format; p.keep_with_next=True
        p.space_before=Pt(20 if course else 0); p.space_after=Pt(10 if course else 0)
        if course and level==1:
            borders=elem('w:pBdr'); borders.append(elem('w:top',w_val='single',w_sz=4,w_space=12,w_color='000000'))
            s.element.get_or_add_pPr().append(borders)
    for name in ['Caption','Image Caption','Table Caption','Footnote Text']:
        s=doc.styles[name]; size=9 if name=='Footnote Text' else 10.5 if course else 16
        font_style(s,cfg['body'],cfg['latin'],size)
        p=s.paragraph_format; p.first_line_indent=Pt(0); p.line_spacing=Pt(size*1.4)
        p.space_before=Pt(4); p.space_after=Pt(9)
    s=doc.styles['Appendix Label']; font_style(s,cfg['hei'],cfg['latin'],16)
    p=s.paragraph_format; p.first_line_indent=Pt(0); p.space_after=Pt(cfg['line']); p.keep_with_next=True
    for name in ['Verbatim Char','Hyperlink']:
        if name in doc.styles: doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    if not course:
        for name in ['Normal','Body Text','First Paragraph','Compact','Heading 1','Heading 2','Heading 3','Heading 4']:
            doc.styles[name].element.get_or_add_rPr().append(elem('w:spacing',w_val=-5))
    font_style(doc.styles['Footer'],cfg['folio'],cfg['folio'],10 if course else 14)
    configure_section(doc.sections[0],profile)
    doc.settings.odd_and_even_pages_header_footer = not course
    for even in ([False] if course else [False,True]):
        footer=doc.sections[0].even_page_footer if even else doc.sections[0].footer
        p=footer.paragraphs[0]; p.style=doc.styles['Footer']; p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=Pt(14)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if course else WD_ALIGN_PARAGRAPH.LEFT if even else WD_ALIGN_PARAGRAPH.RIGHT
        if not course:
            if even: p.paragraph_format.left_indent=Pt(14)
            else: p.paragraph_format.right_indent=Pt(14)
            p.add_run('— ')
        field(p,'PAGE')
        if not course: p.add_run(' —')
        for r in p.runs:
            r.font.name=cfg['folio']; r.font.size=Pt(10 if course else 14)
    settings=doc.settings.element
    mathpr=settings.find(qn('m:mathPr'))
    if mathpr is None: mathpr=elem('m:mathPr'); settings.append(mathpr)
    old=mathpr.find(qn('m:mathFont'))
    if old is not None: mathpr.remove(old)
    mathpr.append(elem('m:mathFont',m_val=cfg['word_math']))
    settings.append(elem('w:updateFields',w_val='true'))
    doc.save(path)


def polish_docx(path, profile, expected_math, orientation):
    doc=Document(path); cfg=STYLES[profile]
    # Pandoc 3.9 emits literal & alignment markers in OMML eqArr rows.
    # Encode them as native alignment properties; otherwise LO prints ampersands.
    for marker in doc._element.xpath('//m:eqArr//m:r[m:t="&"]'):
        rpr=marker.find(qn('m:rPr'))
        if rpr is None: rpr=elem('m:rPr'); marker.insert(0,rpr)
        rpr.append(elem('m:aln',m_val='1'))
        marker.find(qn('m:t')).text=''
    # Native Word sections keep wide tables editable; PDF uses the standard's rotation.
    final=doc._element.body.sectPr
    base=deepcopy(final)
    for p in list(doc.paragraphs):
        if p.text in ('RECTO_PORTRAIT','RECTO_LANDSCAPE'):
            end=deepcopy(base)
            p._p.get_or_add_pPr().append(end)
            for child in list(p._p):
                if child.tag!=qn('w:pPr'): p._p.remove(child)
        elif p.text == 'RECTO_PAGEBREAK':
            p.clear(); p.add_run().add_break(WD_BREAK.PAGE)
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
        elif p.text == 'RECTO_APPENDIX_END':
            p._p.getparent().remove(p._p)
    # Pandoc does not preserve arbitrary keyword delimiters; pass section state directly.
    orientation=[x=='landscape' for x in orientation]
    for i,section in enumerate(doc.sections):
        configure_section(section,profile,orientation[i] if i<len(orientation) else False)
        typ=section._sectPr.find(qn('w:type'))
        if typ is None: section._sectPr.insert(0,elem('w:type',w_val='nextPage'))
        else: typ.set(qn('w:val'),'nextPage')
    for p in doc.paragraphs:
        if p._p.xpath('.//m:oMathPara'):
            pf=p.paragraph_format; pf.first_line_indent=Pt(0); pf.space_before=Pt(14); pf.space_after=Pt(14)
            pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
            previous=p._p.getprevious()
            if previous is not None and previous.tag==qn('w:p'):
                previous.get_or_add_pPr().append(elem('w:keepNext'))
        text_runs=[r for r in p.runs if r.text.strip()]
        if text_runs and all(r.bold for r in text_runs): p.paragraph_format.keep_with_next=True
    for style_name,prefix in [('Image Caption','图'),('Table Caption','表')]:
        n=0
        for p in doc.paragraphs:
            if p.style.name==style_name:
                n+=1
                if not re.match(r'^[图表]\s*\d+',p.text):
                    r=p.add_run(f'{prefix} {n}　'); p._p.insert(1,r._r)
    # Determine page width at each table from section boundaries, working backwards.
    next_width=doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin
    widths={}
    for child in reversed(list(doc._element.body)):
        sect=child.find('w:pPr/w:sectPr',NS)
        if sect is not None:
            size=sect.find('w:pgSz',NS); margin=sect.find('w:pgMar',NS)
            next_width=(int(size.get(qn('w:w')))-int(margin.get(qn('w:left')))-int(margin.get(qn('w:right'))))*635
        if child.tag==qn('w:tbl'): widths[child]=next_width
    for table in doc.tables:
        width=widths[table._tbl]; table.autofit=False
        columns=list(table.columns); values=[c.width or width/len(columns) for c in columns]
        if profile=='gbt9704':
            def cell_width(cell):
                content=''.join(cell._tc.xpath('.//w:t/text() | .//m:t/text()'))
                return sum(1 if ord(ch)>255 else .5 for ch in content)
            values=[max(cell_width(row.cells[i]) for row in table.rows)+2.2 for i in range(len(columns))]
        total=sum(values)
        for c,v in zip(columns,values): c.width=int(width*v/total)
        props=table._tbl.tblPr
        borders=props.find(qn('w:tblBorders'))
        if borders is not None: props.remove(borders)
        borders=elem('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            value='nil' if profile=='reference' and side in ('left','right','insideV') else 'single'
            borders.append(elem('w:'+side,w_val=value,w_sz=4,w_color='000000'))
        props.append(borders)
        margins=props.find(qn('w:tblCellMar'))
        if margins is not None: props.remove(margins)
        margins=elem('w:tblCellMar')
        for side,v in [('top',140),('bottom',140),('left',120),('right',120)]: margins.append(elem('w:'+side,w_w=v,w_type='dxa'))
        props.append(margins)
        for ri,row in enumerate(table.rows):
            if ri==0: row._tr.get_or_add_trPr().append(elem('w:tblHeader'))
            row._tr.get_or_add_trPr().append(elem('w:cantSplit'))
            for ci,cell in enumerate(row.cells):
                cell.width=columns[min(ci,len(columns)-1)].width
                cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    pf=p.paragraph_format; pf.first_line_indent=Pt(0); pf.space_before=Pt(0); pf.space_after=Pt(0)
                    size=10.5 if profile=='reference' else 16
                    pf.line_spacing=Pt(size*1.55)
                    for r in p.runs:
                        r.font.size=Pt(size)
                        fonts=r._r.get_or_add_rPr().find(qn('w:rFonts'))
                        if fonts is None: fonts=elem('w:rFonts'); r._r.get_or_add_rPr().append(fonts)
                        fonts.set(qn('w:eastAsia'),cfg['hei'] if ri==0 and profile=='gbt9704' else cfg['body'])
                        r.bold=ri==0 and profile=='reference'
    doc.save(path)
    with ZipFile(path) as z:
        math_count=sum(len(etree.fromstring(z.read(n)).xpath('//m:oMath',namespaces=NS)) for n in z.namelist() if n.startswith('word/') and n.endswith('.xml'))
        if math_count!=expected_math: raise ValueError(f'Math coverage: {math_count} native OMML equations, expected {expected_math}')
    return math_count


def walk(value):
    if isinstance(value,dict):
        yield value
        for child in value.values(): yield from walk(child)
    elif isinstance(value,list):
        for child in value: yield from walk(child)


def copy_typst_fonts(families, destination):
    patterns = {
        'FandolSong': 'fandol/FandolSong-*.otf',
        'FandolHei': 'fandol/FandolHei-*.otf',
        'FandolKai': 'fandol/FandolKai-*.otf',
        'FandolFang R': 'fandol/FandolFang-*.otf',
        'Liberation Serif': 'liberation/LiberationSerif-*.ttf',
        'New Computer Modern Math': 'newcm/NewCMMath-*.otf',
    }
    for family in families & patterns.keys():
        for font in FONT_DIR.glob(patterns[family]):
            # Keep each selected family intact with its original license and attribution.
            files = [font, *(p for p in font.parent.iterdir() if p.is_file() and p.suffix not in {'.otf', '.ttf'})]
            for path in files:
                target = destination / path.relative_to(FONT_DIR)
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, target)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source',type=Path)
    parser.add_argument('--style',choices=STYLES,default='reference')
    parser.add_argument('--formats',default='typ,pdf')
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    formats=set(args.formats.split(','))
    if not formats or formats-{'typ','pdf','docx'}: parser.error('formats must be typ,pdf,docx')
    source=args.source.resolve(); out=args.output.resolve(); out.mkdir(parents=True,exist_ok=True)
    runtime=check(args.style)
    if not runtime['ready']:
        raise SystemExit(json.dumps(runtime,ensure_ascii=False,indent=2))
    fonts=runtime['fonts']
    STYLES[args.style].update(fonts, title_bold=args.style=='gbt9704' and fonts['title']=='FandolSong')
    if runtime['substitutions']:
        print('Font substitutions: '+json.dumps(runtime['substitutions'],ensure_ascii=False),file=sys.stderr)
    ast=json.loads(run(['pandoc',str(source),'-f','markdown','-t','json']))
    math_count=0; orientations=['portrait']
    seen_appendix=False
    top_appendices=0
    for block in ast['blocks']:
        is_appendix=block.get('t')=='Div' and 'appendix' in block['c'][0][1]
        if seen_appendix and not is_appendix: raise ValueError('Appendices must be top-level blocks at the end of the document')
        if is_appendix: seen_appendix=True; top_appendices+=1
    all_appendices=sum(n.get('t')=='Div' and 'appendix' in n['c'][0][1] for n in walk(ast))
    if all_appendices!=top_appendices: raise ValueError('Nested appendix blocks are not supported')
    for node in walk(ast):
        if node.get('t')=='Math': math_count+=1
        if node.get('t')=='RawBlock' or node.get('t')=='RawInline':
            raise ValueError('Raw markup is not portable. Use semantic Markdown and the documented Divs.')
        if node.get('t')=='Div' and not set(node['c'][0][1]) & {'appendix','pagebreak'}:
            raise ValueError('Unsupported Div: use appendix or pagebreak')
        if node.get('t')=='Div' and 'appendix' in node['c'][0][1]:
            if args.style!='gbt9704': raise ValueError('Appendix Divs require gbt9704 style')
            orientations.append('landscape' if 'landscape' in node['c'][0][1] else 'portrait')
        if node.get('t')=='Image':
            src=node['c'][2][0]
            if '://' in src or src.startswith('data:'): raise ValueError('Use a local image file, not remote URLs')
            image=(source.parent/src).resolve()
            if not image.is_file(): raise FileNotFoundError(image)
            dest=image if image.parent==out/'media' else out/'media'/(sha256(image.read_bytes()).hexdigest()[:8]+'-'+image.name)
            dest.parent.mkdir(exist_ok=True)
            if image!=dest: shutil.copy2(image,dest)
            node['c'][2][0]=str(dest.relative_to(out))
    font_config='#let recto-fonts = ('+', '.join(key+': '+json.dumps(value) for key,value in fonts.items())+')\n'
    ast['meta']['recto-preamble']={'t':'MetaBlocks','c':[{'t':'RawBlock','c':['typst',font_config+(ROOT/'assets/styles'/f'{args.style}.typ').read_text()]}]}
    if args.style=='reference': ast['meta']['recto-course']={'t':'MetaBool','c':True}
    env=dict(os.environ,RECTO_STYLE=args.style)
    payload=json.dumps(ast).encode()
    common=['pandoc','-f','json','--lua-filter',str(ROOT/'scripts/blocks.lua'),'--resource-path',str(out),'--table-caption-position=below']
    outputs=[]
    if formats & {'typ','pdf'}:
        typst_ast=deepcopy(ast)
        for node in walk(typst_ast):
            if node.get('t')=='Image': node['c'][2][0]='../'+node['c'][2][0]
        with tempfile.TemporaryDirectory(prefix='.recto-',dir=out) as temporary:
            target=Path(temporary)/'document.typ'
            run(common+['-t','typst','--standalone','--template',str(ROOT/'assets/document.typst'),'-o',str(target)],
                input=json.dumps(typst_ast).encode(),env=env,cwd=out)
            if 'pdf' in formats:
                run(['typst','compile','--root',str(out),'--font-path',str(FONT_DIR),str(target),str(out/'document.pdf')],cwd=out)
                outputs.append('document.pdf')
            if 'typ' in formats:
                typst_dir=out/'typst'; typst_dir.mkdir(exist_ok=True)
                copy_typst_fonts({name for role,name in fonts.items() if role!='word_math' and (math_count or role!='math')},typst_dir/'fonts')
                font_arg=' --font-path typst/fonts' if (typst_dir/'fonts').is_dir() else ''
                target.write_text('// From the delivery folder: typst compile --root .'+font_arg+' typst/document.typ document.pdf\n'
                                  '// System fonts are not bundled; selected fonts are listed in recto-fonts below.\n'+target.read_text())
                target.replace(typst_dir/'document.typ')
                outputs.append('typst/document.typ')
    if 'docx' in formats:
        with tempfile.TemporaryDirectory(prefix='recto-word-') as tmp:
            ref=Path(tmp)/'reference.docx'; make_reference(ref,args.style)
            target=out/'document.docx'
            run(common+['-t','docx','--reference-doc',str(ref),'-o',str(target)],input=payload,env=env,cwd=out)
            native_math=polish_docx(target,args.style,math_count,orientations)
            outputs.append(target.name)
    normalized=deepcopy(ast)
    for key in ['recto-preamble','recto-course']: normalized['meta'].pop(key,None)
    (out/'content.md').write_bytes(run(['pandoc','-f','json','-t','markdown','--standalone','--wrap=none'],input=json.dumps(normalized).encode()))
    report=dict(style=args.style,source_sha256=sha256(source.read_bytes()).hexdigest(),outputs=outputs,
                fonts=fonts,font_substitutions=runtime['substitutions'],versions=runtime['versions'],
                source_math_count=math_count,native_word_math_count=native_math if 'docx' in formats else None,
                word_visual_qa='pending' if 'docx' in formats else None,
                word_landscape='native landscape sections; PDF uses GB/T rotated portrait pages' if 'landscape' in orientations else None)
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__': main()
